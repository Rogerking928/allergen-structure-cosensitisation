"""每個圖表面板一張分頁的原始數據（Source_data.xlsx）。

數字一律從畫圖用的同一批輸出檔與同一套篩選取，不另算：圖改了這裡就跟著改。
"""
import json
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font as XFont, PatternFill, Alignment

import figs
import structure_cosens as S

OUT = figs.OUT


def _name(x):
    from common import display_name
    return display_name(x)


def _names(df):
    for c in ("a", "b", "component", "left_out"):
        if c in df:
            df[c] = df[c].map(_name)
    return df


def panels():
    """回傳 [(分頁名, 說明, DataFrame)]，順序照圖表編號。"""
    s = figs.summ()
    dd = json.load(open(OUT / "discrimination.json"))
    out = []

    # ---------------- Figure 1
    rows = [("Pairs estimated", s["pairs_pooled"], None, None),
            ("With an AlphaFold model for both", s["pairs_with_structure"],
             s["pairs_pooled"] - s["pairs_with_structure"], "no model for one member"),
            ("Analysed pairs", s["pairs_core"], s["pairs_with_structure"] - s["pairs_core"],
             "same species, sub-chain, < 50 aa, unannotated"),
            ("Homologous", s["pairs_homologous"], s["pairs_nonhomologous"], "non-homologous")]
    out.append(("Fig 1B", "Figure 1B. Derivation of the analysed pairs.",
                pd.DataFrame(rows, columns=["Step", "Pairs", "Removed at this step",
                                            "Reason removed"])))

    d, meta, comps = figs.load()
    rows = []
    for chip in figs.F.CHIP_ORDER:
        g = d[d.Chip_Type == chip]
        x = g[comps].replace(-1, np.nan)
        keep = [c for c in comps if x[c].notna().sum() >= 0.8 * len(g)]
        breadth = 100 * (x[keep] >= 0.3).sum(axis=1) / len(keep)
        h, e = np.histogram(breadth.clip(upper=59.9), bins=np.arange(0, 61, 2))   # 同 figs.py
        for lo, hi, n in zip(e[:-1], e[1:], h):
            rows.append([figs.F.CHIP_LABEL[chip], len(g), len(keep), lo, hi, int(n),
                         100 * n / len(breadth)])
    out.append(("Fig 1C", "Figure 1C. Sensitisation breadth: percentage of a platform's "
                          "components to which each patient is positive (2-point bins).",
                pd.DataFrame(rows, columns=["Platform", "Patients", "Components measured",
                                            "Bin from (%)", "Bin to (%)", "Patients in bin",
                                            "Patients in bin (% of platform)"])))

    p = figs.load_pairs()
    out.append(("Fig 1D", "Figure 1D. Unadjusted and breadth-adjusted log odds ratio, "
                          "all estimated pairs.",
                _names(p[["a", "b", "beta_crude", "beta", "se"]].copy()).rename(columns={
                    "a": "Component A", "b": "Component B",
                    "beta_crude": "Unadjusted log OR", "beta": "Adjusted log OR",
                    "se": "SE of adjusted log OR"})))

    # ---------------- Figure 2
    pp, c, ss = figs._pairs()
    hom, non = c[c.homologous], c[~c.homologous]
    g = pd.concat([ss.assign(Group="Same species"), hom.assign(Group="Homologous"),
                   non.assign(Group="Non-homologous")])
    g = _names(g[["Group", "a", "b", "beta", "se"]].copy())
    g["Odds ratio"] = np.exp(g.beta)
    out.append(("Fig 2A", "Figure 2A. Co-sensitisation by pair group.",
                g.rename(columns={"a": "Component A", "b": "Component B",
                                  "beta": "Adjusted log OR", "se": "SE"})))
    h = _names(hom[["a", "b", "tm", "seqid80", "beta", "se"]].copy())
    h["Odds ratio"] = np.exp(h.beta)
    h["Meets Codex criterion (>= 35%)"] = h.seqid80 >= 35
    out.append(("Fig 2B", "Figure 2B. Homologous pairs against TM-score (point area "
                          "proportional to 1/SE).",
                h.rename(columns={"a": "Component A", "b": "Component B", "tm": "TM-score",
                                  "seqid80": "Best 80-aa window identity (%)",
                                  "beta": "Adjusted log OR", "se": "SE"})))
    tb = pd.DataFrame(s["tm_bins"])
    tb["median_OR"] = np.exp(tb.median_logOR)
    out.append(("Fig 2B bands", "Figure 2B. Red segments: median adjusted log OR by TM-score "
                                "band.", tb))
    n_ = _names(non[["a", "b", "tm", "beta"]].copy())
    n_["Odds ratio"] = np.exp(n_.beta)
    out.append(("Fig 2C", "Figure 2C. Non-homologous pairs (shown as density).",
                n_.rename(columns={"a": "Component A", "b": "Component B", "tm": "TM-score",
                                   "beta": "Adjusted log OR"})))
    out.append(("Fig 2D", "Figure 2D. Median odds ratio by TM-score band, homologous pairs.",
                pd.read_csv(OUT / "tm_bands_ci.csv").rename(columns={
                    "band": "TM-score band", "n": "Pairs", "median_or": "Median OR",
                    "lo": "95% CI lower", "hi": "95% CI upper", "pct_z3": "% with z > 3",
                    "example": "Example pair", "example_or": "Example OR"})))

    # ---------------- Figure 3
    cc = figs.core()
    a3 = _names(cc[["a", "b", "seqid80", "tm", "beta", "z", "homologous"]].copy())
    a3["Structural threshold (TM >= 0.5)"] = a3.tm >= S.TM_HI
    a3["Codex criterion (>= 35% over 80 aa)"] = a3.seqid80 >= S.SEQ_HI
    a3["Strongly co-sensitised (z > 3)"] = a3.z > 3
    out.append(("Fig 3A", "Figure 3A. All analysed pairs by sequence and structural "
                          "similarity (identity 0 = no detectable local similarity, "
                          "plotted at 'none').",
                a3.rename(columns={"a": "Component A", "b": "Component B",
                                   "seqid80": "Best 80-aa window identity (%)",
                                   "tm": "TM-score", "beta": "Adjusted log OR",
                                   "homologous": "Homologous"})))
    out.append(("Fig 3B", "Figure 3B. Pairs and strongly co-sensitised pairs in each cell "
                          "of the two criteria (also the quadrant labels of A).",
                pd.DataFrame(dd["reclassification"]).rename(columns={
                    "structural": "TM-score >= 0.5", "codex": "Codex criterion",
                    "n": "Pairs", "strong": "Strongly co-sensitised",
                    "pct_strong": "% strongly co-sensitised", "median_or": "Median OR"})))
    q = pd.read_csv(OUT / "quadrant_struct_only.csv")
    top = q[q.z > 3].head(14).copy()
    top["Odds ratio"] = np.exp(top.beta)
    top["95% CI lower"] = np.exp(top.beta - 1.96 * top.se)
    top["95% CI upper"] = np.exp(top.beta + 1.96 * top.se)
    out.append(("Fig 3C", "Figure 3C. The 14 pairs with the highest z among those below the "
                          "sequence criterion with TM-score >= 0.5 (all listed in Table S3).",
                _names(top[["a", "b", "tm", "seqid80", "Odds ratio", "95% CI lower",
                            "95% CI upper", "z"]]).rename(columns={
                    "a": "Component A", "b": "Component B", "tm": "TM-score",
                    "seqid80": "Best 80-aa window identity (%)"})))
    roc = []
    for col, lab in [("tm", "TM-score"), ("seqid80", "80-aa identity")]:
        r = pd.read_csv(OUT / f"roc_{col}.csv")
        roc.append(r.rename(columns={"fpr": f"{lab}: 1 - specificity",
                                     "tpr": f"{lab}: sensitivity"}).reset_index(drop=True))
    out.append(("Fig 3D ROC", "Figure 3D. ROC curves for strong co-sensitisation.",
                pd.concat(roc, axis=1)))
    auc = pd.DataFrame(dd["auc_all"]).rename(columns={"lo": "95% CI lower",
                                                      "hi": "95% CI upper"})
    auc = auc[auc.predictor.isin(["tm", "seqid80"])].replace(
        {"predictor": {"tm": "TM-score", "seqid80": "80-aa identity"}})
    pts = pd.DataFrame(dd["points_all"]).replace(
        {"criterion": {"codex": "Codex criterion", "struct": "TM-score >= 0.5"}})
    cd = json.load(open(OUT / "coef_diff.json"))
    sm = {(r["analysis"], r["term"]): r for r in s["robust"]}
    jt = sm[("joint model", "TM (rank, SD)")]
    js = sm[("joint model", "sequence identity (rank, SD)")]
    joint = pd.DataFrame([
        ["Joint model, homologous pairs: TM-score coefficient per SD of rank",
         cd["b_tm"], *cd["b_tm_ci"]],
        ["Joint model, homologous pairs: identity coefficient per SD of rank",
         cd["b_seq"], *cd["b_seq_ci"]],
        ["Difference, structure minus sequence", cd["diff"], *cd["diff_ci"]]],
        columns=["Quantity", "Estimate", "95% CI lower", "95% CI upper"])
    assert abs(jt["estimate"] - cd["b_tm"]) < 1e-6 and abs(js["estimate"] - cd["b_seq"]) < 1e-6
    out.append(("Fig 3D AUC", "Figure 3D. Areas under the curve, binary criteria (points) "
                              "and joint model (text in panel).",
                [auc, pts, joint]))

    # ---------------- Figure 4
    strong = cc[cc.z > 3].copy()
    fam = figs.families.table()
    fam = fam[fam.status == "ok"].drop_duplicates("component").set_index("component").family
    e = strong[["a", "b", "z", "beta", "homologous"]].copy()
    e["Family A"], e["Family B"] = e.a.map(fam), e.b.map(fam)
    out.append(("Fig 4A edges", "Figure 4A. Network edges: strongly co-sensitised pairs "
                                "(z > 3).",
                _names(e).rename(columns={"a": "Component A", "b": "Component B",
                                          "beta": "Adjusted log OR",
                                          "homologous": "Homologous"})))
    nodes = sorted(set(strong.a) | set(strong.b))
    out.append(("Fig 4A nodes", "Figure 4A. Network nodes and colouring family.",
                pd.DataFrame({"Component": [_name(n) for n in nodes],
                              "Family": [fam.get(n) for n in nodes],
                              "Coloured family": [fam.get(n) if fam.get(n) in figs.FAM_COLORS
                                                  else "other family" for n in nodes]})))
    fams = list(figs.FAM_COLORS)
    rows = []
    for fa in fams:
        for fb in fams:
            g = cc[((cc.fam_a == fa) & (cc.fam_b == fb)) | ((cc.fam_a == fb) & (cc.fam_b == fa))]
            rows.append([fa, fb, len(g), g.beta.median() if len(g) >= 3 else None])
    out.append(("Fig 4B", "Figure 4B. Median adjusted log OR for each pair of families "
                          "(blank where fewer than three pairs).",
                pd.DataFrame(rows, columns=["Family (row)", "Family (column)", "Pairs",
                                            "Median adjusted log OR"])))
    ov = json.load(open(OUT / "overlay.json"))
    out.append(("Fig 4C-D", "Figure 4C, D. Superposed AlphaFold models (TM-align).",
                pd.DataFrame([{"Panel": pnl, "Pair": k.replace("|", " - ").replace("_", " "),
                               "TM-score": v["tm"], "80-aa identity (%)": v["seqid80"],
                               "Odds ratio": v["odds_ratio"], "z": v["z"],
                               "Aligned residues": v["aligned"], "RMSD (A)": v["rmsd"]}
                              for pnl, (k, v) in zip("CD", ov.items())])))

    # ---------------- 補充圖
    ci = json.load(open(OUT / "rho_ci.json"))
    lab = {"homologous": "Homologous pairs, main analysis", "isac_v1": "ISAC v1 only",
           "isac_v2": "ISAC v2 only", "alex": "ALEX2 only",
           "different_order": "Source species in different orders",
           "plddt70": "Both AlphaFold models pLDDT >= 70",
           "below_codex": "Below the sequence criterion",
           "nonhomologous": "Non-homologous pairs, main analysis",
           "nonhomologous_nofrag": "Excluding fragment sequences",
           "nonhomologous_curated": "Curated family labels instead of Pfam/SCOP"}
    out.append(("Fig S1A", "Supplementary Figure S1A. Spearman rho between TM-score and "
                           "co-sensitisation, sensitivity analyses.",
                pd.DataFrame([[lab[k], ci[k]["n"], ci[k]["rho"], ci[k]["lo"], ci[k]["hi"]]
                              for k in lab],
                             columns=["Analysis", "Pairs", "rho", "95% CI lower",
                                      "95% CI upper"])))
    out.append(("Fig S1B", "Supplementary Figure S1B. Leaving out one homology group.",
                _names(pd.read_csv(OUT / "robust_leave_one_group.csv")).rename(columns={
                    "left_out": "Group left out", "pairs_removed": "Pairs removed",
                    "rho_tm": "rho (TM-score)", "rho_seq": "rho (80-aa identity)"})))
    cp = pd.read_csv(OUT / "component_plddt.csv")
    out.append(("Fig S2A", "Supplementary Figure S2A. Mean pLDDT of each component's model.",
                _names(cp.copy()).rename(columns={"component": "Component",
                                                  "plddt": "Mean pLDDT"})))
    lut = cp.set_index("component").plddt
    hh = cc[cc.homologous].copy()
    hh["Lower pLDDT of the two"] = np.minimum(hh.a.map(lut), hh.b.map(lut))
    out.append(("Fig S2B", "Supplementary Figure S2B. TM-score against the lower pLDDT, "
                           "homologous pairs.",
                _names(hh[["a", "b", "Lower pLDDT of the two", "tm"]]).rename(columns={
                    "a": "Component A", "b": "Component B", "tm": "TM-score"})))
    b = json.load(open(OUT / "bands.json"))
    nh, nn = np.load(OUT / "null_homologous.npy"), np.load(OUT / "null_nonhomologous.npy")
    m = max(len(nh), len(nn))
    pad = lambda v: np.concatenate([v, np.full(m - len(v), np.nan)])
    out.append(("Fig S3", f"Supplementary Figure S3. Spearman rho under permuted allergen "
                          f"labels. Observed: homologous {b['obs_homologous']:.4f}, "
                          f"non-homologous {b['obs_nonhomologous']:.4f}.",
                pd.DataFrame({"Permutation": np.arange(1, m + 1),
                              "A: homologous pairs": pad(nh),
                              "B: non-homologous pairs": pad(nn)})))
    bc = pd.read_csv(OUT / "cosens_by_chip.csv")
    bc = bc[bc.method == "mle"]
    rows = []
    for chip in figs.F.CHIP_ORDER:
        mm = cc[["a", "b", "tm", "homologous"]].merge(
            bc[bc.chip == chip][["a", "b", "beta"]], on=["a", "b"])
        mm.insert(0, "Platform", figs.F.CHIP_LABEL[chip])
        rows.append(mm)
    out.append(("Fig S4", "Supplementary Figure S4. Adjusted log OR estimated on each "
                          "platform alone, against TM-score.",
                _names(pd.concat(rows)).rename(columns={
                    "a": "Component A", "b": "Component B", "tm": "TM-score",
                    "homologous": "Homologous", "beta": "Adjusted log OR (this platform)"})))
    return out


def write(path, tables):
    """tables：[(分頁名, 標題, DataFrame)]，含 book.py 的表格；照傳入順序寫。"""
    XF, XB = XFont(name="Arial", size=10), XFont(name="Arial", size=10, bold=True)
    TITLE = XFont(name="Arial", size=11, bold=True)
    FILL = PatternFill("solid", start_color="E8ECF0")
    wb = Workbook(); wb.remove(wb.active)
    toc = wb.create_sheet("Contents")
    toc["A1"] = "Source data for all tables and figures"; toc["A1"].font = TITLE
    toc.cell(3, 1, "Sheet").font = XB; toc.cell(3, 2, "Content").font = XB
    r = 4
    for name, title, dfs in tables:
        assert len(name) <= 31, name
        ws = wb.create_sheet(name)
        ws["A1"] = title; ws["A1"].font = TITLE
        row0 = 3
        for df in dfs if isinstance(dfs, list) else [dfs]:
            for j, cname in enumerate(df.columns, 1):
                cell = ws.cell(row0, j, str(cname)); cell.font = XB; cell.fill = FILL
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for i, row in enumerate(df.itertuples(index=False), row0 + 1):
                for j, v in enumerate(row, 1):
                    v = v.item() if hasattr(v, "item") else v
                    if isinstance(v, float) and np.isnan(v):
                        v = None
                    ws.cell(i, j, v).font = XF
            for j, cname in enumerate(df.columns, 1):
                w = max([len(str(cname))] + [len(str(v)) for v in df.iloc[:, j - 1].head(200)])
                col = ws.cell(1, j).column_letter
                ws.column_dimensions[col].width = max(
                    ws.column_dimensions[col].width or 0, min(max(w + 2, 10), 42))
            row0 += len(df) + 3
        toc.cell(r, 1, name).font = XF
        toc.cell(r, 2, title).font = XF
        r += 1
    toc.column_dimensions["A"].width = 16
    toc.column_dimensions["B"].width = 110
    wb.save(path)
