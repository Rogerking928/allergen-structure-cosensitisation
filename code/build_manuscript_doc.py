"""Render manuscript.md as a Word file in the house manuscript format.

The specification was measured from Life_Disruption_Manuscript.docx rather than
recalled: Letter, one-inch margins on all four sides, Times New Roman 12 pt,
double spaced, justified, no first-line indent. Major headings (Abstract,
Introduction, Methods, Results, Discussion, Declaration, the generative AI
declaration and References) are 14 pt bold; numbered subheadings such as
"2.1. Study Population" are 12 pt bold. Reference entries are left aligned with
a half-inch hanging indent. In the abstract only, the structured label at the
start of each paragraph is bold and the rest of the paragraph is not.

  python3 build_manuscript_doc.py [manuscript.md] [Manuscript.docx]
"""
import os
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
BIG = {"Abstract", "Introduction", "Methods", "Results", "Discussion",
       "Declaration", "References",
       "Declaration of generative AI and AI-assisted technologies in the "
       "manuscript preparation process"}
# a journal package may add top-level headings of its own (Springer asks for
# "Statements and Declarations"); they arrive through the environment so the
# house default is unchanged for every other manuscript
BIG |= {h for h in os.environ.get("HOUSE_EXTRA_BIG", "").split("|") if h}

# 小節標題：舊稿用 2.1. 編號；JACI 不編號，改以「§ 」標記、輸出時拿掉
SUB = re.compile(r"^(\d+\.\d+\.\s|§\s)")
REFLINE = re.compile(r"^\d+\.\s+[A-ZÅÄÖÜÉ]")
ABSTRACT_LABEL = re.compile(r"^(Background|Objective|Introduction|Objectives|Methods|Results|Discussion|"
                            r"Conclusions|Registration):")


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "manuscript.md")
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "Manuscript.docx")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11.0)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)

    def para(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first=0.0, left=0.0):
        p = doc.add_paragraph()
        p.alignment = align
        pf = p.paragraph_format
        pf.line_spacing = 2.0
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.first_line_indent = Inches(first)
        pf.left_indent = Inches(left)
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        return p

    lines = [ln.strip() for ln in src.read_text().split("\n")]
    in_abstract = False
    first_line = True
    for ln in lines:
        if not ln:
            continue
        if ln in BIG:
            para(ln, size=14, bold=True)
            in_abstract = (ln == "Abstract")
            first_line = False
            continue
        if first_line:                       # the title carries the same size
            para(ln, bold=True)              # as the body but is bold
            first_line = False
            continue
        if SUB.match(ln):
            para(ln[2:] if ln.startswith("§ ") else ln, bold=True)
            continue
        if REFLINE.match(ln):
            para(ln, align=WD_ALIGN_PARAGRAPH.LEFT, first=-0.5, left=0.5)
            continue
        m = ABSTRACT_LABEL.match(ln) if in_abstract else None
        if m:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing = 2.0
            pf.space_after = Pt(0); pf.space_before = Pt(0)
            for txt, bold in ((m.group(0), True), (ln[m.end():], False)):
                r = p.add_run(txt)
                r.font.name = FONT
                r.font.size = Pt(12)
                r.font.bold = bold
                r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            continue
        para(ln)

    doc.save(out)
    # Passages changed in a revision arrive wrapped in the marker pair and are
    # coloured red where that helper exists; this manuscript is a first draft,
    # so its absence is not an error.
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "analysis"))
    try:
        from redmark import mark_red
        mark_red(out)
    except ImportError:
        pass
    doc = Document(out)
    n_par = len(doc.paragraphs)
    words = sum(len(p.text.split()) for p in doc.paragraphs)
    print(f"wrote {out}  ({n_par} paragraphs, {words:,} words)")


if __name__ == "__main__":
    main()
